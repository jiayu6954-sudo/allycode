#!/usr/bin/env python3
"""Framework-neutral XLSX plain-table tool. No network or shell execution."""
import argparse
from copy import copy
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from zipfile import ZipFile
import xml.etree.ElementTree as ET

MAX_BYTES = 100 * 1024 * 1024
MAX_CELLS = 200_000
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def preflight(path):
    if path.suffix.lower() != '.xlsx':
        raise ValueError('仅支持普通 .xlsx 文件')
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if sum(x.file_size for x in entries) > MAX_BYTES:
            raise ValueError('工作簿解压大小超过 100 MiB')
        blocked = ('xl/drawings/', 'xl/charts/', 'xl/pivot', 'xl/slicer',
                   'xl/externalLinks/', 'xl/embeddings/', 'xl/activeX/',
                   'xl/ctrlProps/', 'xl/threadedComments/', 'xl/persons/',
                   'xl/comments', 'xl/richData/', '_xmlsignatures/', 'customXml/')
        for entry in entries:
            name = entry.filename
            if name.startswith(blocked) or 'vbaProject' in name:
                raise ValueError('发现不支持保真处理的对象：' + name)
            if name.endswith('.xml'):
                data = archive.read(name)
                if b'<!DOCTYPE' in data or b'<!ENTITY' in data:
                    raise ValueError('不支持包含 DTD/实体声明的工作簿')
                root = ET.fromstring(data)
                # WPS/Excel extension objects can be lost by a generic writer.
                if any(x.tag.rsplit('}', 1)[-1] == 'extLst' for x in root.iter()):
                    raise ValueError('发现扩展节点，需使用能保留扩展的原生工具：' + name)
                if name == 'xl/sharedStrings.xml' or name.startswith('xl/worksheets/'):
                    for x in root.iter():
                        if x.tag in (NS + 'si', NS + 'is') and x.find(NS + 'r') is not None:
                            raise ValueError('发现富文本；请使用原生工具去色，避免丢失字符样式语义')


def load(path, template=False):
    from openpyxl import load_workbook
    if not template:
        preflight(path)
    else:
        # Template is read-only: no round-trip, so its harmless extensions need not block.
        if path.suffix.lower() != '.xlsx':
            raise ValueError('模板必须为 .xlsx')
        with ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_BYTES:
                raise ValueError('模板超过大小限制')
    import warnings
    with warnings.catch_warnings():
        if template:
            warnings.simplefilter('ignore', UserWarning)
        wb = load_workbook(path, data_only=False, keep_links=True)
    for ws in wb:
        if ws.max_row * ws.max_column > MAX_CELLS:
            raise ValueError('使用区域过大：' + ws.title)
        if ws.protection.sheet and not template:
            raise ValueError('工作表受保护，不能自动取消保护：' + ws.title)
    return wb


def snapshot(ws):
    return {c.coordinate: (c.value, c.data_type, c.number_format)
            for row in ws for c in row if c.value is not None}


def inspect_workbook(path):
    wb = load(Path(path).expanduser().resolve(), template=True)
    result = {'ok': True, 'sheets': []}
    for ws in wb:
        sample = [[c.value for c in row] for row in ws.iter_rows(
            min_row=1, max_row=min(8, ws.max_row), max_col=min(6, ws.max_column))]
        result['sheets'].append({'name': ws.title, 'state': ws.sheet_state,
            'range': ws.calculate_dimension(), 'tables': {t.name: t.ref for t in ws.tables.values()},
            'formula_count': sum(c.data_type == 'f' for row in ws for c in row), 'sample': sample})
    wb.close()
    return result


def options(req):
    if not isinstance(req, dict):
        raise ValueError('请求必须是 JSON 对象')
    schema = json.loads((Path(__file__).resolve().parent.parent / 'assets' / 'tool.schema.json').read_text(encoding='utf-8'))
    unknown = set(req) - set(schema['properties'])
    if unknown:
        raise ValueError('未知参数：' + ', '.join(sorted(unknown)))
    for key in ('input', 'output'):
        if not isinstance(req.get(key), str) or not req[key]:
            raise ValueError('缺少路径参数：' + key)
    for key in ('acknowledge_crop', 'overwrite', 'header_bold', 'header_border'):
        if key in req and type(req[key]) is not bool:
            raise ValueError(key + ' 必须为布尔值')
    for key in ('template_header_row', 'template_body_row'):
        if key in req and (type(req[key]) is not int or req[key] < 1):
            raise ValueError(key + ' 必须为正整数')
    for key in ('template', 'template_sheet', 'font'):
        if key in req and (not isinstance(req[key], str) or not req[key]):
            raise ValueError(key + ' 必须为非空字符串')
    if 'font_size' in req and (type(req['font_size']) not in (float, int) or not 6 <= req['font_size'] <= 32):
        raise ValueError('font_size 必须在 6 到 32 之间')
    mode = req.get('mode', 'preserve')
    if mode not in ('preserve', 'compact') or req.get('layout', 'preserve') not in ('preserve', 'readable'):
        raise ValueError('无效的 mode 或 layout')
    for key in ('headers', 'ranges'):
        if key in req and not isinstance(req[key], dict):
            raise ValueError(key + ' 必须为对象')
    if mode == 'preserve' and ('ranges' in req or req.get('acknowledge_crop')):
        raise ValueError('preserve 不接受裁剪参数')
    if mode == 'compact' and ('headers' in req or req.get('acknowledge_crop') is not True):
        raise ValueError('compact 需要 acknowledge_crop=true，且表头由 ranges 第一行决定')
    return mode


def template_style(req):
    style = dict(font='宋体', font_size=11, header_bold=True, header_border=True,
                 header_horizontal='center', header_vertical='top')
    if req.get('template'):
        wb = load(Path(req['template']).expanduser().resolve(), template=True)
        ws = wb[req.get('template_sheet', wb.sheetnames[0])]
        h = ws.cell(req.get('template_header_row', 1), 1)
        b = ws.cell(req.get('template_body_row', 2), 1)
        style.update(font=b.font.name or '宋体', font_size=b.font.sz or 11,
                     header_bold=bool(h.font.bold),
                     header_border=any(getattr(h.border, edge).style for edge in ('left', 'right', 'top', 'bottom') if getattr(h.border, edge) is not None),
                     header_horizontal=h.alignment.horizontal or 'center',
                     header_vertical=h.alignment.vertical or 'top')
        wb.close()
    for key in ('font', 'font_size', 'header_bold', 'header_border'):
        if key in req:
            style[key] = req[key]
    return style


def compact_workbook(wb, req):
    from openpyxl import Workbook
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    from openpyxl.worksheet.cell_range import CellRange
    ranges = req.get('ranges', {})
    if set(ranges) != set(wb.sheetnames):
        raise ValueError('compact ranges 必须精确覆盖所有工作表')
    if len(wb.defined_names):
        raise ValueError('compact 不支持定义名称；仅去色请使用 preserve')
    new = Workbook()
    new.remove(new.active)
    new.epoch = wb.epoch
    for ws in wb:
        if ws.merged_cells.ranges or len(ws.defined_names) or any(c.data_type == 'f' for row in ws for c in row):
            raise ValueError('compact 不支持公式、合并单元格或定义名称：' + ws.title)
        for table in ws.tables.values():
            if any(col.calculatedColumnFormula or col.totalsRowFormula for col in table.tableColumns):
                raise ValueError('compact 不支持结构化表公式：' + ws.title)
        ref = ranges[ws.title]
        if not isinstance(ref, str) or not re.fullmatch(r'[A-Z]+[1-9][0-9]*:[A-Z]+[1-9][0-9]*', ref):
            raise ValueError('无效范围：' + str(ref))
        a, b, c, d = range_boundaries(ref)
        if a > c or b > d or c > ws.max_column or d > ws.max_row:
            raise ValueError('范围超出原始使用区域：' + ref)
        target = new.create_sheet(ws.title)
        target.sheet_state = ws.sheet_state
        for row in ws.iter_rows(min_row=b, max_row=d, min_col=a, max_col=c):
            for cell in row:
                dest = target.cell(cell.row-b+1, cell.column-a+1)
                dest.value, dest.data_type = cell.value, cell.data_type
                dest.number_format = cell.number_format
                if cell.hyperlink:
                    if cell.hyperlink.location or (cell.hyperlink.target or '').startswith('#'):
                        raise ValueError('compact 不支持内部超链接：' + ws.title)
                    dest.hyperlink = copy(cell.hyperlink)
                    dest.hyperlink.ref = dest.coordinate
        for col in range(a, c+1):
            old = ws.column_dimensions.get(get_column_letter(col))
            if old:
                target.column_dimensions[get_column_letter(col-a+1)].width = old.width
        for row in range(b, d+1):
            target.row_dimensions[row-b+1].height = ws.row_dimensions[row].height
        for dv in ws.data_validations.dataValidation:
            for formula in (dv.formula1, dv.formula2):
                if formula and not (re.fullmatch(r'"[^\"]*"', str(formula)) or re.fullmatch(r'-?[0-9]+(?:\.[0-9]+)?', str(formula))):
                    raise ValueError('compact 不支持引用型验证规则：' + ws.title)
            shifted = []
            for part in dv.ranges.ranges:
                if part.min_col >= a and part.max_col <= c and part.min_row >= b and part.max_row <= d:
                    cr = CellRange(str(part)); cr.shift(col_shift=1-a, row_shift=1-b)
                    shifted.append(str(cr))
            if shifted:
                clone = copy(dv); clone.sqref = ' '.join(shifted)
                target.add_data_validation(clone)
        target.freeze_panes = 'A2'
    return new


def apply_style(ws, style, header_rows, readable):
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils.cell import get_column_letter
    body = Font(name=style['font'], size=style['font_size'], color='FF000000')
    header = Font(name=style['font'], size=style['font_size'], color='FF000000', bold=style['header_bold'])
    empty, no_border = PatternFill(fill_type=None), Border()
    side = Side(style='thin', color='FF000000')
    header_border = Border(left=side, right=side, top=side, bottom=side) if style['header_border'] else no_border
    ws.sheet_properties.tabColor = None
    ws.sheet_view.showGridLines = True
    ws.conditional_formatting = __import__('openpyxl').formatting.formatting.ConditionalFormattingList()
    for table in ws.tables.values():
        table.tableStyleInfo = None
        for obj in [table, *table.tableColumns]:
            for attr in ('headerRowDxfId', 'dataDxfId', 'totalsRowDxfId', 'headerRowBorderDxfId', 'tableBorderDxfId', 'totalsRowBorderDxfId'):
                if hasattr(obj, attr):
                    setattr(obj, attr, None)
    for dimension in [*ws.row_dimensions.values(), *ws.column_dimensions.values()]:
        dimension.font, dimension.fill, dimension.border = body, empty, no_border
    for row in ws:
        lines = 1
        for cell in row:
            is_header = cell.row in header_rows
            cell.font, cell.fill = header if is_header else body, empty
            cell.border = header_border if is_header else no_border
            align = copy(cell.alignment)
            if is_header:
                align.horizontal, align.vertical = style['header_horizontal'], style['header_vertical']
            if readable:
                align.wrap_text, align.vertical = True, 'center'
                align.shrink_to_fit = False
            cell.alignment = align
            if readable:
                col = get_column_letter(cell.column)
                width = ws.column_dimensions[col].width or 13
                width = max(8, min(width, 60))
                ws.column_dimensions[col].width = width
                text = str(cell.value) if cell.value is not None else ''
                n = sum(max(1, math.ceil(sum(2 if ord(ch)>255 else 1 for ch in line)/max(1,width-2))) for line in text.split('\n'))
                lines = max(lines, n)
        if readable:
            ws.row_dimensions[row[0].row].height = min(409, max(20, lines*(style['font_size']*1.4)+6))


def format_workbook(req):
    mode = options(req)
    source = Path(req['input']).expanduser().resolve()
    output = Path(req['output']).expanduser().resolve()
    if output.suffix.lower() != '.xlsx':
        raise ValueError('输出必须为 .xlsx')
    for protected in [source] + ([Path(req['template']).expanduser().resolve()] if req.get('template') else []):
        if output == protected or (output.exists() and protected.exists() and os.path.samefile(output, protected)):
            raise ValueError('禁止覆盖输入或模板文件')
    if output.exists() and not req.get('overwrite', False):
        raise ValueError('输出已存在；使用新路径，或明确设置 overwrite=true')
    wb = load(source)
    names = wb.sheetnames[:]
    if mode == 'compact':
        original = wb
        wb = compact_workbook(original, req)
        original.close()
    before = {ws.title: snapshot(ws) for ws in wb}
    style = template_style(req)
    headers = req.get('headers', {})
    if set(headers) - set(names):
        raise ValueError('headers 含不存在的工作表')
    for ws in wb:
        if mode == 'compact':
            rows = {1}
        elif ws.title in headers:
            h = headers[ws.title]
            if type(h) is not int or not 1 <= h <= ws.max_row:
                raise ValueError('表头行超出范围：' + ws.title)
            rows = {h}
        elif ws.tables:
            from openpyxl.utils.cell import range_boundaries
            rows = {range_boundaries(t.ref)[1] for t in ws.tables.values() if t.headerRowCount}
        else:
            rows = {next((c.row for row in ws for c in row if c.value is not None), 1)}
        apply_style(ws, style, rows, req.get('layout', 'preserve') == 'readable')
    if wb.calculation:
        wb.calculation.fullCalcOnLoad = True
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='.plain-excel-', suffix='.xlsx', dir=output.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        from openpyxl import load_workbook
        check = load_workbook(temporary, data_only=False)
        if check.sheetnames != names:
            raise ValueError('校验失败：工作表顺序改变')
        for ws in check:
            if snapshot(ws) != before[ws.title]:
                raise ValueError('校验失败：数据/类型/数字格式变化：' + ws.title)
            for row in ws:
                for cell in row:
                    if cell.value is not None and (cell.fill.patternType is not None or cell.font.color is None or cell.font.color.type != 'rgb' or cell.font.color.rgb != 'FF000000'):
                        raise ValueError('校验失败：发现非黑字或填充色')
        sheets = [{'name': w.title, 'rows': w.max_row, 'columns': w.max_column} for w in check]
        check.close()
        wb.close()
        if req.get('overwrite', False):
            os.replace(temporary, output)
        else:
            os.link(temporary, output)  # Atomic, exclusive publication; same filesystem.
            os.unlink(temporary)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return {'ok': True, 'output': str(output), 'mode': mode, 'sheets': sheets,
            'checks': {'cell_values_types_formats': True, 'black_text_no_fill': True},
            'warnings': ['未进行 Excel/WPS 可视化检查。', '公式缓存可能清空；有公式时需用电子表格程序重算。']}


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    sub.add_parser('inspect').add_argument('--input', required=True)
    sub.add_parser('run').add_argument('--request', required=True)
    args = parser.parse_args()
    try:
        if args.command == 'inspect':
            result = inspect_workbook(args.input)
        else:
            result = format_workbook(json.loads(Path(args.request).read_text(encoding='utf-8-sig')))
        print(json.dumps(result, ensure_ascii=False, default=str))
        return 0
    except Exception as exc:
        print(json.dumps({'ok': False, 'error': str(exc)}, ensure_ascii=False))
        return 2


if __name__ == '__main__':
    sys.exit(main())
