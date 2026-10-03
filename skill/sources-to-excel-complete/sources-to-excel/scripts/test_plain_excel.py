"""Portable regression checks. Run: python scripts/test_plain_excel.py"""
from datetime import datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from plain_excel import format_workbook


class PlainExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.xlsx'
        self.output = self.root / 'output.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def request(self, **kw):
        return dict(input=str(self.source), output=str(self.output), **kw)

    def test_preserves_formulas_types_formats_and_validation(self):
        wb = Workbook(); ws = wb.active
        ws.append(['编号', '金额', '日期', '公式', '状态'])
        ws.append(['00123', 0, datetime(2026, 9, 16), '=B2+7', '未联系'])
        ws['B2'].number_format = '0.00'
        ws['C2'].number_format = 'yyyy-mm-dd'
        for row in ws:
            for c in row:
                c.fill = PatternFill('solid', fgColor='FF0000')
                c.font = Font(color='0000FF')
        validation = DataValidation(type='list', formula1='"未联系,已联系"')
        ws.add_data_validation(validation); validation.add('E2:E20')
        t = Table(displayName='Customers', ref='A1:E2')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(t); wb.save(self.source)
        source_bytes = self.source.read_bytes()
        result = format_workbook(self.request())
        self.assertTrue(result['ok'])
        self.assertEqual(self.source.read_bytes(), source_bytes)
        out = load_workbook(self.output); ws = out.active
        self.assertEqual(ws['A2'].value, '00123')
        self.assertEqual(ws['B2'].value, 0)
        self.assertEqual(ws['B2'].number_format, '0.00')
        self.assertEqual(ws['C2'].value, datetime(2026, 9, 16))
        self.assertEqual(ws['D2'].value, '=B2+7')
        self.assertEqual(ws['D2'].data_type, 'f')
        self.assertEqual(str(ws.data_validations.dataValidation[0].sqref), 'E2:E20')
        self.assertIn('Customers', ws.tables)
        self.assertIsNone(ws.tables['Customers'].tableStyleInfo)
        self.assertEqual(ws['B2'].font.color.rgb, 'FF000000')
        self.assertIsNone(ws['B2'].fill.patternType)
        out.close()

    def test_compact_with_offset_and_leading_zero(self):
        wb = Workbook(); ws = wb.active
        ws['A1']='可删除标题'; ws['B4']='编号'; ws['C4']='电话'; ws['D4']='状态'
        ws['B5']='0001'; ws['C5']='0999-1234567'; ws['D5']='未联系'
        dv=DataValidation(type='list',formula1='"未联系,已联系"')
        ws.add_data_validation(dv); dv.add('D5:D5'); wb.save(self.source)
        format_workbook(self.request(mode='compact',ranges={'Sheet':'B4:D5'},acknowledge_crop=True))
        out=load_workbook(self.output); ws=out.active
        self.assertEqual(list(ws.values), [('编号','电话','状态'),('0001','0999-1234567','未联系')])
        self.assertEqual(str(ws.data_validations.dataValidation[0].sqref),'C2')
        self.assertEqual(ws.freeze_panes,'A2'); out.close()

    def test_compact_rejects_formula_before_output(self):
        wb=Workbook(); wb.active['A1']='表头'; wb.active['A2']='=1+1'; wb.save(self.source)
        with self.assertRaisesRegex(ValueError,'公式'):
            format_workbook(self.request(mode='compact',ranges={'Sheet':'A1:A2'},acknowledge_crop=True))
        self.assertFalse(self.output.exists())

    def test_existing_output_unchanged_on_failure(self):
        wb=Workbook(); wb.active['A1']='表头'; wb.save(self.source)
        self.output.write_bytes(b'original-output')
        with self.assertRaisesRegex(ValueError,'已存在'):
            format_workbook(self.request())
        self.assertEqual(self.output.read_bytes(),b'original-output')
        with self.assertRaisesRegex(ValueError,'覆盖输入'):
            format_workbook({'input':str(self.source),'output':str(self.source),'overwrite':True})

    def test_template_content_never_copied(self):
        wb=Workbook(); wb.active.append(['客户']); wb.active.append(['真实内容']); wb.save(self.source)
        template=self.root/'template.xlsx'
        t=Workbook(); t.active['A1']='不应复制'; t.active['A2']='模板业务数据'
        t.active['A2'].font=Font(name='Arial',size=12); t.save(template)
        format_workbook(self.request(template=str(template),header_bold=False,header_border=False))
        out=load_workbook(self.output); ws=out.active
        self.assertEqual(list(ws.values),[('客户',),('真实内容',)])
        self.assertEqual(ws['A2'].font.name,'Arial')
        self.assertFalse(ws['A1'].font.bold)
        self.assertTrue(ws['A1'].border.top is None or ws['A1'].border.top.style is None)
        out.close()


if __name__ == '__main__':
    unittest.main()
