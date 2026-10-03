"""Bounded local formula validation and verified cached-value publication, not a formula evaluator."""
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
from uuid import uuid4
from decimal import Decimal
from zipfile import ZipFile
import xml.etree.ElementTree as E
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries

FUNCTIONS = set('SUM AVERAGE COUNT COUNTA MIN MAX ROUND ROUNDUP ROUNDDOWN ABS IF IFERROR AND OR NOT SUMIF SUMIFS COUNTIF COUNTIFS AVERAGEIF AVERAGEIFS SUMPRODUCT SUBTOTAL VLOOKUP HLOOKUP INDEX MATCH XLOOKUP DATE YEAR MONTH DAY LEN LEFT RIGHT MID TRIM CONCATENATE TEXT'.split())


def formula_value(value):
    if not isinstance(value, dict) or not isinstance(value.get('formula'), str) or not value['formula'].startswith('=') or len(value['formula']) > 4000:
        raise ValueError('公式字段须提供 formula（以 = 开头，最长 4000 字符）')
    if not isinstance(value.get('explanation'), str) or not value['explanation'].strip():
        raise ValueError('公式必须说明业务计算口径 explanation')
    kind = value.get('resultType', 'decimal')
    expected = value.get('expected')
    if kind not in {'decimal', 'integer', 'text', 'boolean'} or expected is None:
        raise ValueError('公式须提供独立核算的 expected 和 resultType（decimal/integer/text/boolean）')
    if kind in {'decimal', 'integer'}:
        if isinstance(expected, (bool, float)) or not Decimal(str(expected)).is_finite():
            raise ValueError('公式 expected 数字须用有限十进制字符串或整数')
        if kind == 'integer' and Decimal(str(expected)) != Decimal(str(expected)).to_integral_value():
            raise ValueError('整数公式 expected 不能含小数')
    elif not isinstance(expected, str if kind == 'text' else bool):
        raise ValueError('公式 expected 类型不符')
    tolerance = Decimal(str(value.get('tolerance', '0.000000001')))
    if not tolerance.is_finite() or not Decimal(0) <= tolerance <= Decimal('0.01'):
        raise ValueError('公式误差容忍范围必须在 0 至 0.01 之间，金额建议 0.000001 或更严格')
    fmt = value.get('numberFormat', '0.00' if kind == 'decimal' else 'General')
    if not isinstance(fmt, str) or len(fmt) > 100:
        raise ValueError('公式 numberFormat 无效')
    expression = '=' + ''.join('_xlfn.XLOOKUP(' if t.type == 'FUNC' and t.subtype == 'OPEN' and t.value.upper() == 'XLOOKUP(' else t.value for t in Tokenizer(value['formula']).items)
    return expression, fmt, {**value, 'formula': expression, 'resultType': kind, 'tolerance': str(tolerance)}


def validate_formulas(wb, formulas):
    if len(formulas) > 10000:
        raise ValueError('单批最多 10000 个公式，请按工作表/实体分批')
    graph, total = {}, 0
    sheets = {name.casefold(): name for name in wb.sheetnames}
    for spec in formulas:
        deps = set()
        for token in Tokenizer(spec['formula']).items:
            if token.type == 'FUNC' and token.subtype == 'OPEN':
                name = token.value[:-1].upper().removeprefix('_XLFN.')
                if name not in FUNCTIONS:
                    raise ValueError('暂不支持或不允许此函数：' + name)
            if token.type == 'OPERAND' and token.subtype == 'RANGE':
                ref = token.value
                match = re.fullmatch(r"(?:(?:'((?:[^']|'')+)'|([^'!]+))!)?(\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?)", ref)
                if not match:
                    raise ValueError('仅支持当前工作簿明确的 A1 单元格/有限区域引用，禁止外链、名称和整列范围：' + ref)
                name = (match[1].replace("''", "'") if match[1] else match[2]) or spec['sheet']
                if name.casefold() not in sheets:
                    raise ValueError('公式引用不存在的工作表：' + name)
                name = sheets[name.casefold()]
                bounds = range_boundaries(match[3])
                c1, r1, c2, r2 = bounds
                if c2 < c1 or r2 < r1 or c2 > wb[name].max_column or r2 > wb[name].max_row:
                    raise ValueError('公式引用超出已生成数据范围：' + ref)
                count = (c2-c1+1)*(r2-r1+1)
                total += count
                if count > 200000 or total > 2000000:
                    raise ValueError('公式依赖范围过大，请拆分批次')
                for row in wb[name].iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                    for cell in row:
                        if cell.data_type == 'f':
                            deps.add((name, cell.coordinate))
            if token.type in {'ARRAY'} or (token.type == 'OPERAND' and token.subtype == 'ERROR'):
                raise ValueError('不支持数组常量或错误引用，需先修复公式')
        graph[(spec['sheet'], spec['cell'])] = deps
    # Iterative topological removal avoids recursion limits on long formula chains.
    pending = {key: set(deps) for key, deps in graph.items()}
    reverse = {}
    for key, deps in pending.items():
        for dep in deps:
            reverse.setdefault(dep, set()).add(key)
    ready = [key for key, deps in pending.items() if not deps]
    while ready:
        key = ready.pop()
        pending.pop(key, None)
        for follower in reverse.get(key, ()):
            if follower in pending:
                pending[follower].discard(key)
                if not pending[follower]:
                    ready.append(follower)
    if pending:
        raise ValueError('检测到循环公式依赖：' + str(list(pending)[:5]))


def finalize(r):
    original, calculated, output = map(Path, (r['original'], r['calculated'], r['output']))
    if hashlib.sha256(original.read_bytes()).hexdigest() != r['sha256']:
        raise ValueError('重算期间暂存工作簿改变')
    audit = json.loads(Path(str(original)+'.audit.json').read_text(encoding='utf-8'))
    specs = audit['formulaChecks']
    for source in audit['manifest']['files']:
        if source.get('sha256') and hashlib.sha256(Path(source['path']).read_bytes()).hexdigest() != source['sha256']:
            raise ValueError('重算期间来源资料改变')
    values = load_workbook(calculated, data_only=True, keep_links=False)
    formulas = load_workbook(original, data_only=False, keep_links=False)
    checked = []
    try:
        for spec in specs:
            cell = values[spec['sheet']][spec['cell']]
            actual, expected, kind = cell.value, spec['expected'], spec['resultType']
            if actual is None and kind == 'text' and expected == '':
                with ZipFile(calculated) as archive:
                    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    part = 'xl/worksheets/sheet'+str(values.sheetnames.index(spec['sheet'])+1)+'.xml'
                    node = E.fromstring(archive.read(part)).find(".//m:c[@r='"+spec['cell']+"']", ns)
                    if node is not None and node.get('t') == 'str' and node.find('m:v', ns) is not None:
                        actual = ''
            if cell.data_type == 'e' or actual is None:
                raise ValueError(f"公式错误或无计算缓存：{spec['sheet']}!{spec['cell']} = {actual}")
            if kind in {'decimal', 'integer'}:
                valid = isinstance(actual, (int, float)) and not isinstance(actual, bool) and math.isfinite(actual) and abs(Decimal(str(actual))-Decimal(str(expected))) <= Decimal(spec['tolerance'])
            else:
                valid = type(actual) is type(expected) and actual == expected
            if not valid:
                raise ValueError(f"公式与独立核算不一致：{spec['sheet']}!{spec['cell']}，计算={actual}，预期={expected}")
            checked.append({**spec, 'actual': actual, 'status': 'passed'})
    finally:
        values.close()
    # Transfer only verified caches; preserve the original workbook's styles,
    # literal text and formula expressions instead of shipping a Calc rewrite.
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    tag = '{'+ns['m']+'}'
    cache_by_part = {}
    for check in checked:
        index = formulas.sheetnames.index(check['sheet']) + 1
        cache_by_part.setdefault(f'xl/worksheets/sheet{index}.xml', {})[check['cell']] = check
    formulas.close()
    nonce = uuid4().hex
    staged = output.parent / (output.stem+'.'+nonce+'.verified.xlsx')
    staged_audit = output.parent / (output.name+'.'+nonce+'.audit.tmp')
    output.parent.mkdir(parents=True, exist_ok=True)
    linked = []
    try:
        with ZipFile(original) as src, staged.open('xb') as stream, ZipFile(stream, 'w') as dst:
            for item in src.infolist():
                raw = src.read(item.filename)
                if item.filename in cache_by_part:
                    root = E.fromstring(raw)
                    for cell in root.findall('.//m:sheetData/m:row/m:c', ns):
                        check = cache_by_part[item.filename].get(cell.get('r'))
                        if check is None:
                            continue
                        for old in cell.findall('m:v', ns):
                            cell.remove(old)
                        kind, value = check['resultType'], check['actual']
                        cell.set('t', 'str' if kind == 'text' else 'b' if kind == 'boolean' else 'n')
                        E.SubElement(cell, tag+'v').text = ('1' if value else '0') if kind == 'boolean' else str(value)
                    raw = E.tostring(root, encoding='utf-8', xml_declaration=True)
                dst.writestr(item, raw)
        verify = load_workbook(staged, data_only=True)
        expressions = load_workbook(staged, data_only=False)
        try:
            for spec in checked:
                cached = verify[spec['sheet']][spec['cell']].value
                if spec['resultType'] == 'text' and spec['actual'] == '' and cached is None:
                    cached = ''
                if cached != spec['actual'] or expressions[spec['sheet']][spec['cell']].value != spec['formula']:
                    raise ValueError('最终公式/缓存重开校验失败')
        finally:
            verify.close()
            expressions.close()
        result = {**audit['result'], 'output': str(output), 'audit': str(output)+'.audit.json', 'formulas': {'status': 'recalculated_and_verified', 'count': len(checked), 'engine': 'LibreOffice Calc'}}
        audit.update(result=result, formulaChecks=checked)
        audit['request']['output'] = str(output)
        staged_audit.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
        for src, dst in [(staged_audit, Path(result['audit'])), (staged, output)]:
            os.link(src, dst)
            linked.append(dst)
        return result
    except Exception:
        for file in linked:
            file.unlink()
        raise
    finally:
        staged.unlink(missing_ok=True)
        staged_audit.unlink(missing_ok=True)


if __name__ == '__main__':
    print(json.dumps(finalize(json.load(sys.stdin)), ensure_ascii=False))
