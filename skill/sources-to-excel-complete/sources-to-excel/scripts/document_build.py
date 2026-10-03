"""AllyCode Word writer using the user's mandatory Chinese document profile."""
import json
import sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_font(style, east, latin, size, bold=False):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.underline = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key, value in [('eastAsia', east), ('ascii', latin), ('hAnsi', latin), ('cs', latin)]:
        fonts.set(qn('w:' + key), value)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)


def spacing(style, points, indent=0):
    f = style.paragraph_format
    f.line_spacing = Pt(points)
    f.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    f.space_before = Pt(0)
    f.space_after = Pt(0)
    f.first_line_indent = Pt(indent)


def add_page_number(footer, align, p):
    para = footer.paragraphs[0]
    para.alignment = align
    para.paragraph_format.first_line_indent = Pt(0)
    run = para.add_run()
    run.font.name = p['latinFont']
    run.font.size = Pt(p['pageNumberSizePt'])
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), p['pageNumberFont'])
    begin, instruction, end = (OxmlElement('w:' + k) for k in ('fldChar', 'instrText', 'fldChar'))
    begin.set(qn('w:fldCharType'), 'begin')
    instruction.text = ' PAGE '
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, end])


def build(r):
    d, p = Document(), r['profile']
    # The stock template can carry an accent rule on Title and italic Heading 4.
    # Neither belongs to the user's supplied government-document reference.
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for attr, key in [('top_margin','marginTopCm'), ('bottom_margin','marginBottomCm'), ('left_margin','marginLeftCm'), ('right_margin','marginRightCm'), ('header_distance','headerCm'), ('footer_distance','footerCm')]:
        setattr(s, attr, Cm(p[key]))
    set_font(d.styles['Normal'], p['bodyFont'], p['latinFont'], p['bodySizePt'])
    spacing(d.styles['Normal'], p['bodyLinePt'], p['bodySizePt'] * p['firstLineChars'])
    set_font(d.styles['Title'], p['titleFont'], p['latinFont'], p['titleSizePt'])
    spacing(d.styles['Title'], p['titleLinePt'])
    for level in range(1, 5):
        style = d.styles['Heading ' + str(level)]
        set_font(style, p['headingFonts'][level-1], p['latinFont'], p['headingSizePt'], p['headingBold'][level-1])
        spacing(style, p['bodyLinePt'])
        style.paragraph_format.keep_with_next = True
    title = d.add_paragraph(r['title'], 'Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    for _ in range(p['titleBlankLines']):
        blank = d.add_paragraph('')
        blank.paragraph_format.keep_with_next = True
    d.add_paragraph(r['summary'])
    counters = [0, 0, 0, 0]
    def cn(n):
        digits = '零一二三四五六七八九'
        if n < 10: return digits[n]
        if n < 20: return '十' + (digits[n % 10] if n % 10 else '')
        if n < 100: return digits[n // 10] + '十' + (digits[n % 10] if n % 10 else '')
        return str(n)
    for section in r['sections']:
        level = section.get('level', 1)
        counters[level-1] += 1
        for i in range(level, 4): counters[i] = 0
        n = counters[level-1]
        prefix = [cn(n)+'、', '（'+cn(n)+'）', str(n)+'. ', '（'+str(n)+'）'][level-1]
        d.add_heading(prefix + section['heading'], level=level)
        for text in section.get('paragraphs', []): d.add_paragraph(text)
        for data in section.get('tables', []):
            table = d.add_table(rows=1, cols=len(data['headers']))
            table.style = 'Table Grid'
            table.autofit = False
            width = (s.page_width-s.left_margin-s.right_margin)//len(data['headers'])
            for col in table.columns: col.width = width
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell, value in zip(table.rows[0].cells, data['headers']):
                cell.text = str(value)
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7E6E6')
                cell._tc.get_or_add_tcPr().append(shade)
            for values in data['rows']:
                for cell, value in zip(table.add_row().cells, values): cell.text = str(value)
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs: para.paragraph_format.first_line_indent = Pt(0)
            d.add_paragraph()
    if r.get('attachments'):
        d.add_paragraph()
        for i, item in enumerate(r['attachments'], 1):
            lead = '附件：' if i == 1 else ''
            if len(r['attachments']) > 1: lead += str(i) + '. '
            para = d.add_paragraph(lead + item.rstrip('。；;，,'))
            para.paragraph_format.left_indent = Pt(p['bodySizePt'] * 2)
            para.paragraph_format.first_line_indent = Pt(0)
    if r.get('issuer') or r.get('date'):
        d.add_paragraph()
        d.add_paragraph()
        center = int((s.page_width-s.left_margin-s.right_margin) * 0.75)
        for value in [r.get('issuer'), r.get('date')]:
            if value:
                para = d.add_paragraph('\t' + value)
                para.paragraph_format.first_line_indent = Pt(0)
                para.paragraph_format.tab_stops.add_tab_stop(center, WD_TAB_ALIGNMENT.CENTER)
    if r.get('contact'): d.add_paragraph('（' + r['contact'].strip('（）()') + '）')
    d.settings.odd_and_even_pages_header_footer = True
    add_page_number(s.footer, WD_ALIGN_PARAGRAPH.RIGHT, p)
    add_page_number(s.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT, p)
    with open(r['output'], 'xb') as stream: d.save(stream)
    reopened = Document(r['output'])
    assert reopened.paragraphs[0].text == r['title']
    assert len(reopened.tables) == sum(len(s.get('tables', [])) for s in r['sections'])
    print(json.dumps({'output':r['output'], 'paragraphs':len(reopened.paragraphs), 'tables':len(reopened.tables)}, ensure_ascii=False))


if __name__ == '__main__': build(json.load(sys.stdin))
