#!/usr/bin/env python3
"""Portable source inventory and evidence-backed XLSX writer; agent supplies semantics/OCR."""
import argparse
import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from datetime import date, datetime
from decimal import Decimal
from zipfile import ZipFile
import xml.etree.ElementTree as ET

MAX_BYTES = 100 * 1024 * 1024
MAX_CELLS = 200_000
TEXT_EXT = {'.txt', '.md', '.json', '.jsonl', '.log'}
IMAGE_EXT = {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.webp', '.heic'}


def dumps(obj):
    return json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False)


def plain_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def check_archive(path):
    with ZipFile(path) as z:
        if sum(x.file_size for x in z.infolist()) > MAX_BYTES:
            raise ValueError('解压后超过 100 MiB；需分批读取')


def decode(path):
    raw = path.read_bytes()
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            return raw.decode(encoding), encoding
        except UnicodeError:
            pass
    raise ValueError('无法无损解码；请指定编码后重新读取')


def extract(path):
    ext = path.suffix.lower()
    segments, warnings = [], []
    if ext in TEXT_EXT or ext in {'.csv', '.tsv'}:
        text, encoding = decode(path)
        warnings.append('解码方式：' + encoding)
        if ext in {'.csv', '.tsv'}:
            delimiter = '\t' if ext == '.tsv' else ','
            reader = csv.reader(io.StringIO(text, newline=''), delimiter=delimiter)
            for row, values in enumerate(reader, 1):
                segments.append({'locator': f'row:{row}', 'text': dumps(values), 'values': values})
        else:
            for line, value in enumerate(text.splitlines(), 1):
                segments.append({'locator': f'line:{line}', 'text': value})
        return segments, warnings, 'extracted'
    if ext == '.xlsx':
        from openpyxl import load_workbook
        check_archive(path)
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for ws in wb:
                if (ws.max_row or 0) * (ws.max_column or 0) > MAX_CELLS:
                    raise ValueError(f'{ws.title} 范围过大；需分区读取')
                for row in ws:
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = plain_value(cell.value)
                        seg = {'locator': f'{ws.title}!{cell.coordinate}', 'text': str(value),
                               'value': value, 'data_type': cell.data_type,
                               'number_format': cell.number_format, 'sheet_state': ws.sheet_state}
                        if cell.data_type == 'f':
                            seg['cached_value'] = plain_value(cached[ws.title][cell.coordinate].value)
                            warnings.append(f'{seg["locator"]} 为公式，缓存需核实')
                        segments.append(seg)
            with ZipFile(path) as z:
                if any(n.startswith('xl/media/') for n in z.namelist()):
                    warnings.append('存在嵌入图片；需视觉检查其中是否含业务数据')
                    return segments, warnings, 'needs_visual_review'
            return segments, warnings, 'extracted'
        finally:
            wb.close()
            cached.close()
    if ext == '.docx':
        check_archive(path)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with ZipFile(path) as z:
            parts = [n for n in z.namelist() if n == 'word/document.xml' or
                     re.match(r'word/(header\d+|footer\d+|footnotes|endnotes)\.xml$', n)]
            for part in sorted(parts):
                raw = z.read(part)
                if b'<!DOCTYPE' in raw or b'<!ENTITY' in raw:
                    raise ValueError('不支持 XML DTD/实体')
                root = ET.fromstring(raw)
                # Keep paragraph evidence for prose and structured table evidence
                # for cell pairing. They are alternate views, not additive rows.
                for i, para in enumerate(root.findall('.//w:p', ns), 1):
                    text = ''.join(x.text or '' for x in para.findall('.//w:t', ns))
                    if text:
                        style = para.find('w:pPr/w:pStyle', ns)
                        segments.append({'locator': f'{part}:paragraph:{i}', 'text': text,
                                         'style': style.get('{'+ns['w']+'}val') if style is not None else None})
                for ti, table in enumerate(root.findall('.//w:tbl', ns), 1):
                    for ri, row in enumerate(table.findall('w:tr', ns), 1):
                        cells = []
                        for ci, cell in enumerate(row.findall('w:tc', ns), 1):
                            text = '\n'.join(''.join(t.text or '' for t in para.findall('.//w:t', ns))
                                             for para in cell.findall('w:p', ns))
                            span = cell.find('w:tcPr/w:gridSpan', ns)
                            merge = cell.find('w:tcPr/w:vMerge', ns)
                            cells.append({'cell': ci, 'text': text,
                                          'colspan': int(span.get('{'+ns['w']+'}val', '1')) if span is not None else 1,
                                          'vmerge': merge.get('{'+ns['w']+'}val', 'continue') if merge is not None else None})
                        segments.append({'locator': f'{part}:table:{ti}:row:{ri}',
                                         'text': dumps([c['text'] for c in cells]),
                                         'values': [c['text'] for c in cells], 'cells': cells,
                                         'kind': 'table_row'})
            warnings.append('表格行与段落是同一内容的两种证据视图，不得重复计数；合并单元格、嵌套表、文本框和图片需按原件复核')
            return segments, warnings, 'needs_visual_review'
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('PDF 加密，需可读取版本')
        for i, page in enumerate(reader.pages, 1):
            segments.append({'locator': f'page:{i}', 'text': page.extract_text() or ''})
        warnings.append('PDF 提取文本不保证阅读顺序；逐页检查扫描页、混合图片及表格')
        return segments, warnings, 'needs_visual_review'
    if ext in IMAGE_EXT:
        return [], ['需调用宿主视觉或 OCR；此脚本不内置识别模型'], 'needs_ocr'
    return [], ['未内置此格式解析器；使用宿主对应读取工具并记录结果'], 'unsupported'


def scan(request):
    inputs = request.get('inputs')
    if not isinstance(inputs, list) or not inputs:
        raise ValueError('inputs 必须是非空文件/目录路径列表')
    output = Path(request['output']).expanduser().resolve()
    if output.exists():
        raise ValueError('扫描输出已存在；请使用新路径')
    excluded = {Path(p).expanduser().resolve() for p in request.get('exclude', [])}
    excluded.add(output)
    candidates = []
    for value in inputs:
        root = Path(value).expanduser().absolute()
        if root.is_symlink():
            candidates.append((root, 'symlink_skipped'))
        elif root.is_dir():
            def onerror(error):
                candidates.append((Path(error.filename), 'unreadable'))
            for parent, dirs, files in os.walk(root, followlinks=False, onerror=onerror):
                kept = []
                for d in sorted(dirs):
                    p = Path(parent) / d
                    if p.resolve() in excluded:
                        continue
                    if p.is_symlink():
                        candidates.append((p, 'symlink_skipped'))
                    else:
                        kept.append(d)
                dirs[:] = kept
                candidates.extend((Path(parent) / f, None) for f in sorted(files))
        else:
            candidates.append((root, None))
    files, seen, hashes = [], set(), {}
    for path, status in candidates:
        absolute = str(path.absolute())
        if absolute in seen or path.resolve() in excluded:
            continue
        seen.add(absolute)
        entry = {'id': f'F{len(files)+1:06}', 'path': absolute, 'status': status,
                 'segments': [], 'warnings': []}
        files.append(entry)
        if path.is_symlink() or status:
            entry['status'] = status or 'symlink_skipped'
            entry['warnings'] = ['未跟随符号链接或无法访问目录；需单独处理']
            continue
        try:
            entry['bytes'] = path.stat().st_size
            if entry['bytes'] > MAX_BYTES:
                raise ValueError('文件超过 100 MiB；需分批读取')
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            entry['sha256'] = digest
            if digest in hashes:
                entry.update(status='duplicate', duplicate_of=hashes[digest])
                continue
            hashes[digest] = entry['id']
            entry['segments'], entry['warnings'], entry['status'] = extract(path)
            if path.suffix.lower() == '.docx':
                tables = {}
                for seg in entry['segments']:
                    if seg.get('kind') != 'table_row':
                        continue
                    key = seg['locator'].rsplit(':row:', 1)[0]
                    info = tables.setdefault(key, {'locator': key, 'rows': 0, 'columnCounts': [], 'merged': False, 'firstRow': seg['values']})
                    info['rows'] += 1
                    width = sum(c['colspan'] for c in seg['cells'])
                    if width not in info['columnCounts']:
                        info['columnCounts'].append(width)
                    info['merged'] |= any(c['colspan'] > 1 or c['vmerge'] is not None for c in seg['cells'])
                entry['structure'] = {
                    'tables': list(tables.values()),
                    'headings': [{'locator': s['locator'], 'text': s['text']} for s in entry['segments'] if str(s.get('style', '')).startswith(('Title', 'Heading'))],
                    'note': '这是结构画像，不是语义判定。首行未必是表头；Agent 应结合全文自行确定实体、字段、单位和记录粒度。',
                }
        except Exception as e:
            entry.update(status='unreadable', warnings=[f'{type(e).__name__}: {e}'])
    manifest = {'version': 1, 'files': files,
                'counts': {s: sum(f['status'] == s for f in files) for s in sorted({f['status'] for f in files})}}
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('x', encoding='utf-8') as stream:
        stream.write(dumps(manifest))
    return {'ok': True, 'output': str(output), 'files': len(files), 'counts': manifest['counts']}


def typed(value, kind):
    if value is None:
        return None, 'General', None
    if kind == 'text':
        if not isinstance(value, str):
            raise ValueError('文本字段须传字符串，防止编号丢失前导零或精度')
        if len(value) > 32767:
            raise ValueError('单元格文本超过 Excel 限制；需拆分并保留来源')
        return value, '@', None
    if kind in {'integer', 'decimal'}:
        if isinstance(value, (bool, float)):
            raise ValueError('数字须传十进制字符串或整数，避免二进制浮点输入误差')
        number = Decimal(str(value))
        if not number.is_finite() or (kind == 'integer' and number != number.to_integral_value()):
            raise ValueError('数字无效或整数含小数')
        if len(number.as_tuple().digits) > 15 or number.adjusted() > 14 or number.adjusted() < -307:
            return str(value), '@', '超过 Excel 可靠数字精度，已作为文本保留'
        decimals = min(20, max(0, -number.as_tuple().exponent))
        converted = int(number) if kind == 'integer' else float(number)
        if not math.isfinite(converted) or Decimal(str(converted)) != number:
            return str(value), '@', '无法无损写为 Excel 数字，已作为文本保留'
        fmt = '0.##############E+00' if -number.as_tuple().exponent > 20 else '0' + ('.' + '0' * decimals if decimals else '')
        return converted, fmt, None
    if kind == 'date':
        return date.fromisoformat(value), 'yyyy-mm-dd', None
    if kind == 'boolean' and isinstance(value, bool):
        return value, 'General', None
    raise ValueError('不支持的字段类型或值：' + str(kind))


def verify_evidence(evidence, files):
    if not isinstance(evidence, list) or not evidence:
        raise ValueError('每个非空字段必须有来源证据')
    ids = set()
    for ev in evidence:
        source = files.get(ev.get('file_id'))
        if source is None:
            raise ValueError('来源文件 ID 不存在')
        match = [s for s in source.get('segments', []) if s['locator'] == ev.get('locator')]
        if not ev.get('quote') or not any(ev['quote'] in s['text'] for s in match):
            raise ValueError('证据定位/原文不匹配：' + str(ev))
        ids.add(source['id'])
    return ids


def build(request):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Border, Alignment
    # Trusted sibling application helper, not the user's working directory.
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from excel_formulas import formula_value, validate_formulas
    manifest_path = Path(request['manifest']).resolve()
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    files = {f['id']: f for f in manifest['files']}
    if len(files) != len(manifest['files']):
        raise ValueError('manifest 文件 ID 重复')
    for source in files.values():
        if source.get('sha256'):
            p = Path(source['path'])
            if not p.is_file() or hashlib.sha256(p.read_bytes()).hexdigest() != source['sha256']:
                raise ValueError('源文件已变动或不可读，请重新扫描：' + source['path'])
    coverage = request.get('coverage', {})
    if set(coverage) != set(files):
        raise ValueError('coverage 必须交代每个扫描文件，不得遗漏')
    for fid, item in coverage.items():
        if item.get('status') not in {'included', 'partial', 'duplicate', 'out_of_scope', 'unreadable'}:
            raise ValueError('coverage 状态无效')
        if item['status'] != 'included' and not item.get('reason'):
            raise ValueError('未纳入的文件必须说明原因')
        if item['status'] in {'included', 'partial'} and files[fid]['status'] in {'needs_ocr', 'needs_visual_review', 'unsupported', 'unreadable'} and not item.get('review_note'):
            raise ValueError('需说明视觉/OCR/备用读取完成情况')
    output = Path(request['output']).resolve()
    audit = Path(str(output) + '.audit.json')
    sources = {Path(f['path']).resolve() for f in files.values()} | {manifest_path}
    if output.suffix.lower() != '.xlsx' or output in sources or audit in sources:
        raise ValueError('输出必须为新 .xlsx，不能覆盖来源')
    if output.exists() or audit.exists():
        raise ValueError('输出或审计文件已存在；请使用新路径')
    sheets = request.get('sheets', [])
    if not sheets:
        raise ValueError('必须有至少一张工作表')
    wb = Workbook()
    wb.remove(wb.active)
    used_ids, warnings, expected, names = set(), [], {}, set()
    formula_checks = []
    for spec in sheets:
        name, columns, rows = spec['name'], spec['columns'], spec['rows']
        if not name or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name) or name.lower() in names:
            raise ValueError('工作表名称无效或重复')
        names.add(name.lower())
        keys = [c['key'] for c in columns]
        headers = [c['header'] for c in columns]
        if not keys or len(set(keys)) != len(keys) or len(set(headers)) != len(headers) or not all(isinstance(x, str) and x.strip() for x in keys + headers):
            raise ValueError('字段名和表头必须非空且不重复')
        if len(columns) > 16384 or len(rows) > 1048575 or (len(rows)+1)*len(columns) > MAX_CELLS:
            raise ValueError('工作表过大；按实体或分批拆分')
        ws = wb.create_sheet(name)
        expected[name] = []
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.data_type = 's'
            expected[name].append((1, col, header, 's', 'General'))
        for r, row in enumerate(rows, 2):
            values, evidence = row['values'], row.get('evidence', {})
            if set(values) != set(keys):
                raise ValueError('每行 values 字段必须与 columns 完全一致；缺失值用 null')
            if any(v is None for v in values.values()) and not row.get('issues'):
                raise ValueError('缺失字段须在 issues 说明，不可静默填空')
            for c, column in enumerate(columns, 1):
                key = column['key']
                if values[key] is not None:
                    used_ids |= verify_evidence(evidence.get(key), files)
                is_formula = column.get('type') == 'formula'
                if is_formula:
                    value, fmt, spec = formula_value(values[key])
                    warning = None
                else:
                    value, fmt, warning = typed(values[key], column.get('type', 'text'))
                cell = ws.cell(r, c, value)
                if is_formula:
                    cell.data_type = 'f'
                    formula_checks.append({**spec, 'sheet': name, 'cell': cell.coordinate})
                elif isinstance(value, str):
                    cell.data_type = 's'  # Untrusted strings beginning '=' must never become formulas.
                cell.number_format = fmt
                expected[name].append((r, c, value, cell.data_type, fmt))
                if warning:
                    warnings.append(f'{name}!{cell.coordinate}: {warning}')
        for row in ws:
            for cell in row:
                cell.font = Font(name='宋体', size=11, color='000000', bold=False)
                cell.fill = PatternFill(fill_type=None)
                cell.border = Border()
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for c in range(1, len(columns)+1):
            from openpyxl.utils import get_column_letter
            lengths = [sum(2 if ord(ch) > 127 else 1 for ch in str(ws.cell(r,c).value or '')) for r in range(1, ws.max_row+1)]
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(12, max(lengths)+2))
        for row in ws:
            lines = max((sum(max(1, math.ceil(sum(2 if ord(ch)>127 else 1 for ch in line) / max(1, ws.column_dimensions[cell.column_letter].width-2))) for line in str(cell.value or '').split('\n')) for cell in row), default=1)
            ws.row_dimensions[row[0].row].height = min(409, 16*lines)
        ws.sheet_view.showGridLines = True
    validate_formulas(wb, formula_checks)
    included = {fid for fid, item in coverage.items() if item['status'] in {'included', 'partial'}}
    if used_ids != included:
        raise ValueError('included 文件必须有数据证据；数据证据不得引用排除文件')
    for fid, item in coverage.items():
        if item['status'] == 'duplicate':
            target = item.get('duplicate_of')
            if target not in included:
                raise ValueError('重复文件须指向已纳入文件 duplicate_of')
    output.parent.mkdir(parents=True, exist_ok=True)
    result = {'ok': True, 'output': str(output), 'audit': str(audit),
              'sheets': {s['name']: len(s['rows']) for s in sheets}, 'warnings': warnings,
              'unresolved_files': [fid for fid, i in coverage.items() if i['status'] in {'unreadable', 'partial'}],
              'review_rows': sum(bool(r.get('issues')) for s in sheets for r in s['rows'])}
    result['formulas'] = {'status': 'pending_recalculation' if formula_checks else 'not_applicable', 'count': len(formula_checks)}
    audit_data = {'result': result, 'request': request, 'manifest': manifest, 'formulaChecks': formula_checks}
    linked = []
    with tempfile.TemporaryDirectory(dir=output.parent) as temp:
        staged = Path(temp) / 'output.xlsx'
        wb.save(staged)
        check = load_workbook(staged)
        try:
            for name, cells in expected.items():
                for r, c, value, dtype, fmt in cells:
                    cell = check[name].cell(r, c)
                    actual = cell.value
                    if isinstance(actual, datetime) and isinstance(value, date) and not isinstance(value, datetime):
                        actual = actual.date()
                    # Empty string XLSX cells are blank on reopen; callers should use null.
                    if value == '':
                        value = None
                    if actual != value or cell.number_format != fmt or (value is not None and cell.data_type != dtype):
                        raise ValueError(f'写后校验不一致：{name}!{cell.coordinate}')
                    if cell.fill.patternType or cell.font.color.rgb != '00000000':
                        raise ValueError('样式校验失败')
        finally:
            check.close()
        staged_audit = Path(temp) / 'audit.json'
        staged_audit.write_text(dumps(audit_data), encoding='utf-8')
        try:
            for src, dst in ((staged_audit, audit), (staged, output)):
                os.link(src, dst)
                linked.append(dst)
        except Exception:
            for dst in linked:
                dst.unlink()
            raise
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['scan', 'build'])
    parser.add_argument('--request', required=True)
    args = parser.parse_args()
    try:
        request = json.loads(Path(args.request).read_text(encoding='utf-8'))
        result = (scan if args.command == 'scan' else build)(request)
        print(dumps(result))
    except Exception as e:
        print(dumps({'ok': False, 'error': f'{type(e).__name__}: {e}'}))
        sys.exit(2)


if __name__ == '__main__':
    main()
